"""Fill the cyber deck with a ZeroTrace hackathon pitch. Layout stays."""
import shutil
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SRC = r"c:\Users\msara\Downloads\Blue Black and White Modern Cyber Security Presentation.pptx"
DST = r"c:\Users\msara\OneDrive\Desktop\ZeroTrace\ZeroTrace-Hackathon-Pitch.pptx"


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh.shapes)
        elif sh.has_text_frame:
            yield sh


def paint(shape, lines):
    paras = list(shape.text_frame.paragraphs)
    for i, p in enumerate(paras):
        text = lines[i] if i < len(lines) else ""
        if not p.runs:
            continue
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""


def by_text(slide):
    found = {}
    for sh in walk(slide.shapes):
        key = " ".join(sh.text_frame.text.split())
        found[key] = sh
    return found


def main():
    shutil.copyfile(SRC, DST)
    prs = Presentation(DST)

    # Slide 1
    s = by_text(prs.slides[0])
    paint(s["BUILDING A SECURE DIGITAL FUTURE"], ["SHARE FILES. LEAVE ZERO TRACE."])
    paint(
        s["Protecting digital assets throgh effective security strategies"],
        ["Scan on the device. Export a clean copy. Nothing is uploaded."],
    )

    # Slide 2
    s = by_text(prs.slides[1])
    paint(
        s[
            "On-device app that scans a file, shows privacy risks, and exports a clean copy. No upload, no account, no cloud scan. Working build: Android release APK"
        ],
        [
            "On-device app: scan a file, show the risks, export a clean copy. No account. No cloud scan. Android APK and Windows app are already built."
        ],
    )

    # Slide 4 — the leak
    s = by_text(prs.slides[3])
    paint(s["Cyber Security Challenges"], ["Where the leak hides"])
    paint(
        s[
            "Organizations face a variety of security challenges as technology continues to evolve. Understanding these challenges helps businesses prepare effective protection strategies."
        ],
        ["The risk is inside the file people already share. Not only on the network."],
    )
    paint(s["Common Challenges"], ["Four leaks"])
    paint(s["Data Privacy"], ["Hidden GPS"])
    paint(s["System Vulnerabilities"], ["Faces and plates"])
    paint(s["Third-Party Risks"], ["Keys and IDs"])
    paint(s["Human Error"], ["Cloud tools"])
    paint(
        s["Protecting confidential information from unauthorized access and misuse."],
        ["Location and camera data stay inside a shared photo."],
    )
    paint(
        s["Minimizing software and hardware vulnerabilities with regular updates and maintenance."],
        ["A picture can show who was there and which vehicle."],
    )
    paint(
        s["Managing security risks with vendors, suppliers, and partners."],
        ["API keys, tokens, and ID numbers sit in docs and code."],
    )
    paint(
        s["Preventing accidental security incidents through employee awareness and training."],
        ["Other scanners need the private file uploaded first."],
    )

    # Slide 5 — flow
    s = by_text(prs.slides[4])
    paint(s["Paucek and Lage"], ["ZeroTrace"])
    paint(s["Security Framework"], ["How a scan runs"])
    paint(
        s[
            "A cyber security framework provides structured guidelines for managing and improving organizational security."
        ],
        ["One file stays on the phone. Four steps. No server."],
    )
    paint(s["Governance"], ["1. Select"])
    paint(s["Risk Management"], ["2. Scan"])
    paint(s["Security Policies"], ["3. Review"])
    paint(s["Continuous Monitoring"], ["4. Export"])
    paint(
        s["Defines security roles, policies, and decision-making processes."],
        ["Pick a photo, PDF, Office file, or code."],
    )
    paint(
        s["Identifies, evaluates, and prioritizes potential security risks."],
        ["Built-in packs run offline on the device."],
    )
    paint(
        s["Establishes clear rules and procedures for protecting organizational assets."],
        ["User sees GPS, faces, keys, and QR hits."],
    )
    paint(
        s["Continuously monitors systems and networks for unusual activities or threats."],
        ["Save a clean copy. History stays until you delete it."],
    )

    # Slide 6 — packs
    s = by_text(prs.slides[5])
    paint(s["Paucek and Lage"], ["ZeroTrace"])
    paint(s["Protection Strategies"], ["Built-in packs"])
    paint(
        s[
            "Effective protection strategies combine multiple security controls to reduce cyber risks and improve resilience."
        ],
        ["These four ship inside the app and work with no network."],
    )
    paint(s["Strategies"], ["Offline"])
    paint(
        s["Access Management Encryption Network Security Regular Backups"],
        ["Privacy Essentials", "Metadata Cleaner", "Developer Protection", "QR Protection"],
    )

    # Slide 7 — stack
    s = by_text(prs.slides[6])
    paint(s["Paucek and Lage"], ["ZeroTrace"])
    paint(s["Emerging Technologies"], ["Built on device"])
    paint(
        s[
            "New technologies help organizations improve their ability to detect, prevent, and respond to cyber threats."
        ],
        ["The scan never leaves the phone or the PC."],
    )
    paint(s["Artificial Intelligence"], ["Flutter UI"])
    paint(s["Machine Learning"], ["Rust engine"])
    paint(s["Cloud Security"], ["On-device AI"])
    paint(s["Automation"], ["Signed packs"])
    paint(
        s["Analyzes large amounts of data to identify suspicious activities more"],
        ["Android and Windows app."],
    )
    paint(
        s["Learns from previous attack patterns to improve threat detection over time."],
        ["Local scanner, called from the app."],
    )
    paint(
        s["Protects cloud infrastructure, applications, and data from cyber risks."],
        ["Face, plate, and OCR models run here."],
    )
    paint(
        s["Automatically detects and handles routine security events, reducing response time."],
        ["Each pack is signature-checked before install."],
    )

    # Slide 8 — audience
    s = by_text(prs.slides[7])
    paint(s["Paucek and Lage"], ["ZeroTrace"])
    paint(s["Building Security Awareness"], ["Who it is for"])
    paint(
        s[
            "Employees play a significant role in maintaining cyber security. Regular training and awareness programs help reduce human error while encouraging safe digital practices across the organization."
        ],
        ["People who share a file and will not upload it just to clean it."],
    )
    paint(s["Focus Area"], ["Users"])
    paint(
        s["Password Security Phising Awareness"],
        ["People sharing photos", "Developers with API keys"],
    )
    paint(
        s["Device Protection Safe Onlien Behavior"],
        ["Students and journalists", "Anyone staying offline"],
    )

    # Slide 9 — proof
    s = by_text(prs.slides[8])
    paint(s["Paucek and Lage"], ["ZeroTrace"])
    paint(s["Measuring Success"], ["Already built"])
    paint(s["Security Compliance"], ["Release APK"])
    paint(s["System Availability"], ["Works offline"])
    paint(s["Training Completion"], ["Optional packs"])
    paint(s["Incident Response Time"], ["One-tap wipe"])
    paint(
        s["Measures adherence to internal policies and industry regulations."],
        ["Android release build, about 95 MB."],
    )
    paint(
        s["Ensures business systems remain operational with minimal downtime."],
        ["Built-in packs need no network."],
    )
    paint(
        s["Tracks employee participation in cyber security awareness programs."],
        ["Face, plate, and OCR after one download."],
    )
    paint(
        s["Evaluates how quickly security incidents are detected and resolved."],
        ["Deletes scans, exports, packs, and keys."],
    )
    paint(
        s[
            "Cyber security performance should be monitored regularly to ensure continuous improvement and effective risk management"
        ],
        ["Face, plate, and OCR are not perfect. The user reviews every hit before sharing."],
    )

    # Slide 10
    s = by_text(prs.slides[9])
    paint(s["Paucek and Lage"], ["ZeroTrace"])
    paint(s["www.yourwebsite.com"], ["Scan. Review. Export."])
    paint(s["BUILDING A SECURE DIGITAL FUTURE"], ["SHARE FILES. LEAVE ZERO TRACE."])
    paint(
        s["Protecting digital assets throgh effective security strategies"],
        ["Open a photo. Show the hits. Export the clean copy."],
    )

    prs.save(DST)
    print("saved", DST)


if __name__ == "__main__":
    main()
